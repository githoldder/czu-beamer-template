#!/usr/bin/env python3
"""Convert PDF to PPTX by embedding each page as an image."""

import sys
from pathlib import Path
from pdf2image import convert_from_path
from pptx import Presentation
from pptx.util import Emu, Inches

# 16:9 widescreen in EMUs (13.333" x 7.5")
SLIDE_WIDTH = Emu(12192000)
SLIDE_HEIGHT = Emu(6858000)

def pdf_to_pptx(pdf_path: str, output_path: str, dpi: int = 200):
    pdf_file = Path(pdf_path)
    print(f"Converting {pdf_file.name} ({pdf_file.stat().st_size / 1024 / 1024:.1f} MB)...")
    
    images = convert_from_path(pdf_path, dpi=dpi)
    print(f"Extracted {len(images)} pages")
    
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    
    blank = prs.slide_layouts[6]  # blank layout
    
    tmp_dir = Path("/tmp/pdf2pptx_images")
    tmp_dir.mkdir(exist_ok=True)
    
    for i, img in enumerate(images):
        img_path = tmp_dir / f"slide_{i:03d}.png"
        img.save(img_path, "PNG")
        
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img_path), 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
        print(f"  Slide {i+1}/{len(images)} done")
        img_path.unlink()
    
    tmp_dir.rmdir()
    prs.save(output_path)
    out_file = Path(output_path)
    print(f"Saved to {out_file.name} ({out_file.stat().st_size / 1024 / 1024:.1f} MB)")

if __name__ == "__main__":
    pdf = sys.argv[1] if len(sys.argv) > 1 else "beamersapienza.pdf"
    out = sys.argv[2] if len(sys.argv) > 2 else "beamersapienza.pptx"
    dpi = int(sys.argv[3]) if len(sys.argv) > 3 else 200
    pdf_to_pptx(pdf, out, dpi)
