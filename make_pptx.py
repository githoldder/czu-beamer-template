#!/usr/bin/env python3
"""Create PPTX from pre-rendered page images."""

from pathlib import Path
from pptx import Presentation
from pptx.util import Emu

SLIDE_WIDTH = Emu(12192000)   # 13.333" (16:9)
SLIDE_HEIGHT = Emu(6858000)   # 7.5" (16:9)

def images_to_pptx(img_dir: str, output_path: str):
    img_dir = Path(img_dir)
    images = sorted(img_dir.glob("*.png"))
    print(f"Found {len(images)} images")
    
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    blank = prs.slide_layouts[6]
    
    for i, img_path in enumerate(images):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(img_path), 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
        print(f"  Slide {i+1}/{len(images)}: {img_path.name}")
    
    prs.save(output_path)
    out = Path(output_path)
    print(f"Saved: {out.name} ({out.stat().st_size / 1024 / 1024:.1f} MB)")

images_to_pptx("/tmp/pdf2pptx_pages", "beamersapienza.pptx")
